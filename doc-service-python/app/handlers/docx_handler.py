"""
DOCX Handler — Extract and reconstruct Word documents.
Extracts one segment per SENTENCE so the translation API receives complete,
context-rich sentences rather than individual words or multi-kilobyte paragraphs.
Preserves paragraph layout, tables, headers, footers, and Devanagari font.
"""

import io
import re
from collections import defaultdict
from typing import Dict, List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_UNDERLINE

from app.models import Segment, SegmentMeta

# OOXML Word namespace attributes
_W              = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_W_RFONTS       = f"{{{_W}}}rFonts"
_W_ASCII        = f"{{{_W}}}ascii"
_W_H_ANSI       = f"{{{_W}}}hAnsi"
_W_CS           = f"{{{_W}}}cs"
_W_EAST_ASIA    = f"{{{_W}}}eastAsia"
_W_ASCII_THEME  = f"{{{_W}}}asciiTheme"
_W_H_ANSI_THEME = f"{{{_W}}}hAnsiTheme"
_W_CS_THEME     = f"{{{_W}}}cstheme"
_W_EA_THEME     = f"{{{_W}}}eastAsiaTheme"

_DEVANAGARI_FONT = "Nirmala UI"

# Paragraph identity: (location, table_idx, row_idx, col_idx, para_idx)
_ParaKey = Tuple[Optional[str], Optional[int], Optional[int], Optional[int], Optional[int]]


def _contains_devanagari(text: str) -> bool:
    return any("ऀ" <= ch <= "ॿ" for ch in text)


def _set_devanagari_font(run, font_name: str) -> None:
    """Set all four OOXML font slots and strip theme overrides."""
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(_W_RFONTS)
    if rFonts is None:
        rFonts = rPr.makeelement(_W_RFONTS)
        rPr.insert(0, rFonts)
    rFonts.set(_W_ASCII,     font_name)
    rFonts.set(_W_H_ANSI,    font_name)
    rFonts.set(_W_CS,        font_name)
    rFonts.set(_W_EAST_ASIA, font_name)
    for attr in (_W_ASCII_THEME, _W_H_ANSI_THEME, _W_CS_THEME, _W_EA_THEME):
        if attr in rFonts.attrib:
            del rFonts.attrib[attr]


def _split_sentences(text: str) -> List[str]:
    """Split on Nepali danda (।) or Western sentence terminators (.?!).
    Each returned piece is a complete sentence suitable for the translation API.
    """
    parts = re.split(r'(?<=[।.?!])\s+', text.strip())
    return [p.strip() for p in parts if p.strip()]


def _extract_paragraphs(
    paragraphs,
    location: str,
    segments: List[Segment],
    seg_id: int,
    table_idx: Optional[int] = None,
    row_idx_override: Optional[int] = None,
    col_idx_override: Optional[int] = None,
) -> int:
    """Extract one segment per sentence from each non-empty paragraph.

    Sentence-level extraction gives the API full context per call (~50-200 chars)
    without the hallucination that occurs with kilobyte-sized paragraphs.
    The paragraph_idx is stored so reconstruction can group sentences back together.
    """
    for para_idx, para in enumerate(paragraphs):
        full_text = "".join(run.text for run in para.runs)
        if not full_text.strip():
            continue

        first_run = next((r for r in para.runs if r.text.strip()), None)
        font = first_run.font if first_run else None

        sentences = _split_sentences(full_text)
        if not sentences:
            sentences = [full_text.strip()]

        for sentence in sentences:
            meta = SegmentMeta(
                type="docx_sentence",
                paragraph_idx=para_idx,
                run_idx=None,
                bold=font.bold if font else None,
                italic=font.italic if font else None,
                underline=(
                    font.underline is not None and font.underline != WD_UNDERLINE.NONE
                    if font and font.underline is not None else False
                ),
                font_size=font.size.pt if (font and font.size) else None,
                font_name=font.name if font else None,
                font_color=(
                    str(font.color.rgb)
                    if (font and font.color and font.color.rgb) else None
                ),
                location=location,
                table_idx=table_idx,
                row_idx=row_idx_override,
                col_idx=col_idx_override,
            )
            segments.append(Segment(id=seg_id, text=sentence, meta=meta))
            seg_id += 1

    return seg_id


def extract_docx(file_bytes: bytes) -> List[Segment]:
    """
    Extract all sentences from a DOCX file with formatting and position metadata.
    Handles body paragraphs, tables, headers, and footers.
    """
    doc = Document(io.BytesIO(file_bytes))
    segments: List[Segment] = []
    seg_id = 0

    seg_id = _extract_paragraphs(doc.paragraphs, "body", segments, seg_id)

    for tbl_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                seg_id = _extract_paragraphs(
                    cell.paragraphs, "table", segments, seg_id,
                    table_idx=tbl_idx,
                    row_idx_override=r_idx,
                    col_idx_override=c_idx,
                )

    for section in doc.sections:
        if section.header and section.header.paragraphs:
            seg_id = _extract_paragraphs(
                section.header.paragraphs, "header", segments, seg_id
            )
        if section.footer and section.footer.paragraphs:
            seg_id = _extract_paragraphs(
                section.footer.paragraphs, "footer", segments, seg_id
            )

    return segments


def _apply_translation(run, translated: str) -> None:
    """Set translated text on a run and fix the Devanagari font if needed."""
    original = run.text
    leading  = original[: len(original) - len(original.lstrip())]
    trailing = original[len(original.rstrip()):]
    run.text = leading + translated + trailing
    if _contains_devanagari(translated):
        _set_devanagari_font(run, _DEVANAGARI_FONT)


def reconstruct_docx(
    original_bytes: bytes, translated_segments: List[Segment]
) -> bytes:
    """
    Reconstruct a DOCX file with translated text.

    Groups sentence-level segments back to their source paragraph by the
    (location, table_idx, row_idx, col_idx, paragraph_idx) key, joins the
    translated sentences with a space, then places the result in the first
    run of that paragraph while clearing the remaining runs.
    """
    doc = Document(io.BytesIO(original_bytes))

    # Build para_key → [translated sentence, ...] ordered by segment id
    para_texts: Dict[_ParaKey, List[str]] = defaultdict(list)
    for seg in sorted(translated_segments, key=lambda s: s.id):
        key: _ParaKey = (
            seg.meta.location,
            seg.meta.table_idx,
            seg.meta.row_idx,
            seg.meta.col_idx,
            seg.meta.paragraph_idx,
        )
        para_texts[key].append(seg.text)

    def patch(
        paragraphs,
        location: str,
        table_idx: Optional[int] = None,
        row_idx: Optional[int] = None,
        col_idx: Optional[int] = None,
    ) -> None:
        for para_idx, para in enumerate(paragraphs):
            if not "".join(run.text for run in para.runs).strip():
                continue
            key: _ParaKey = (location, table_idx, row_idx, col_idx, para_idx)
            if key not in para_texts:
                continue
            translated = " ".join(para_texts[key])
            runs_with_text = [r for r in para.runs if r.text.strip()]
            if runs_with_text:
                _apply_translation(runs_with_text[0], translated)
                for run in runs_with_text[1:]:
                    run.text = ""

    patch(doc.paragraphs, "body")

    for tbl_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                patch(cell.paragraphs, "table", tbl_idx, r_idx, c_idx)

    for section in doc.sections:
        if section.header and section.header.paragraphs:
            patch(section.header.paragraphs, "header")
        if section.footer and section.footer.paragraphs:
            patch(section.footer.paragraphs, "footer")

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.read()
